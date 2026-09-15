"""Build the review presentation from Markdown slide content."""
from __future__ import annotations
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "chapters" / "slides.md"
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "docs" / "IHACDP_Presentation.pptx"
VIT_LOGO = ROOT / "docs" / "assets" / "vit-logo.png"

TEAL, DARK, SLATE, LIGHT = RGBColor(0x0D, 0x6E, 0x6E), RGBColor(0x0B, 0x1B, 0x24), \
                           RGBColor(0x2B, 0x45, 0x52), RGBColor(0xF2, 0xF8, 0xF8)
W, H = Inches(13.333), Inches(7.5)

# Slides that read better with a figure beside the text
SLIDE_FIGURES = {
    9: "fig_architecture.png",
    10: "fig_triage_sources.png",
    13: "fig_algorithm_contest.png",
    14: "fig_llm_benchmark.png",
    17: "fig_triage_accuracy.png",
    15: "ui_consultation.png",
}


def parse(text: str):
    slides, cur = [], None
    for line in text.splitlines():
        m = re.match(r"^##\s*Slide\s*(\d+)\s*[::]\s*(.+)$", line.strip(), re.I)
        if m:
            cur = {"n": int(m.group(1)), "title": m.group(2).strip(), "body": []}
            slides.append(cur); continue
        if cur is not None:
            cur["body"].append(line)
    for s in slides:
        s["body"] = "\n".join(s["body"]).strip()
    return slides


def add_bar(slide):
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.09))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()


def title_slide(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
    accent = sl.shapes.add_shape(1, 0, Inches(2.62), W, Inches(0.05))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()

    if VIT_LOGO.exists():
        # dark navy slide, dark blue lockup -- it needs a white plate to read at all
        from PIL import Image
        iw, ih = Image.open(VIT_LOGO).size
        h, pad = 0.58, 0.16
        w = h * iw / ih
        plate = sl.shapes.add_shape(5, Inches(0.62), Inches(0.42),
                                    Inches(w + 2 * pad), Inches(h + 2 * pad))
        plate.fill.solid(); plate.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        plate.line.fill.background()
        sl.shapes.add_picture(str(VIT_LOGO), Inches(0.62 + pad), Inches(0.42 + pad),
                              height=Inches(h))
    else:
        logo = ROOT / "frontend" / "assets" / "logo.png"
        if logo.exists():
            sl.shapes.add_picture(str(logo), Inches(0.62), Inches(0.5), height=Inches(0.95))

    tb = sl.shapes.add_textbox(Inches(0.8), Inches(1.75), W - Inches(1.6), Inches(1.0))
    p = tb.text_frame.paragraphs[0]; p.text = "IHACDP"
    p.font.size = Pt(52); p.font.bold = True; p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub = sl.shapes.add_textbox(Inches(0.8), Inches(2.95), W - Inches(1.6), Inches(0.9))
    p = sub.text_frame.paragraphs[0]
    p.text = "Intelligent Healthcare Analytics and Clinical Decision Support Platform"
    p.font.size = Pt(21); p.font.color.rgb = RGBColor(0x9C, 0xC4, 0xC4)

    body = [ln.strip("-* ").strip() for ln in s["body"].splitlines() if ln.strip().startswith(("-", "*"))]
    tb2 = sl.shapes.add_textbox(Inches(0.8), Inches(4.25), W - Inches(1.6), Inches(2.6))
    tf = tb2.text_frame; tf.word_wrap = True
    for i, line in enumerate(body[:7]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        p.font.size = Pt(15); p.font.color.rgb = RGBColor(0xDD, 0xE8, 0xE8); p.space_after = Pt(7)
    return sl


def content_slide(prs, s, fig=None):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_bar(sl)
    tb = sl.shapes.add_textbox(Inches(0.62), Inches(0.36), W - Inches(1.24), Inches(0.85))
    p = tb.text_frame.paragraphs[0]; p.text = s["title"]
    p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = DARK

    rule = sl.shapes.add_shape(1, Inches(0.62), Inches(1.19), Inches(1.5), Inches(0.045))
    rule.fill.solid(); rule.fill.fore_color.rgb = TEAL; rule.line.fill.background()

    body = s["body"]
    rows = [ln for ln in body.splitlines() if ln.strip().startswith("|")]
    width = W - Inches(1.24)
    if fig and (FIG / fig).exists():
        width = Inches(6.7)

    if len(rows) >= 2:
        _table(sl, rows, Inches(0.62), Inches(1.5), width)
    else:
        bullets = [ln.strip("-*• ").strip() for ln in body.splitlines()
                   if ln.strip().startswith(("-", "*", "•"))]
        if not bullets:
            bullets = [ln.strip() for ln in body.splitlines() if ln.strip()]
        _bullets(sl, bullets, Inches(0.62), Inches(1.5), width)

    if fig and (FIG / fig).exists():
        sl.shapes.add_picture(str(FIG / fig), Inches(7.55), Inches(1.6), width=Inches(5.2))
    return sl


def _bullets(sl, bullets, left, top, width):
    tb = sl.shapes.add_textbox(left, top, width, H - top - Inches(0.6))
    tf = tb.text_frame; tf.word_wrap = True
    size = 18 if len(bullets) <= 6 else (16 if len(bullets) <= 9 else 14)
    for i, b in enumerate(bullets[:12]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        b = re.sub(r"\*\*(.+?)\*\*", r"\1", b)
        p.text = "•  " + b
        p.font.size = Pt(size); p.font.color.rgb = SLATE
        p.space_after = Pt(9 if size >= 16 else 6)


def _table(sl, rows, left, top, width):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(set(x) <= set("-: ") for x in c)]
    if not cells:
        return
    ncol = max(len(c) for c in cells)
    cells = [c + [""] * (ncol - len(c)) for c in cells]
    nrow = min(len(cells), 9)
    shape = sl.shapes.add_table(nrow, ncol, left, top, width,
                                Inches(0.42) * nrow).table
    for r in range(nrow):
        for c in range(ncol):
            cell = shape.cell(r, c)
            cell.text = re.sub(r"\*\*(.+?)\*\*", r"\1", cells[r][c])
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(13 if ncol <= 4 else 11)
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else SLATE
            para.font.bold = r == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if r == 0 else (
                LIGHT if r % 2 else RGBColor(0xFF, 0xFF, 0xFF))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def closing_slide(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
    tb = sl.shapes.add_textbox(Inches(0.8), Inches(2.9), W - Inches(1.6), Inches(1.4))
    p = tb.text_frame.paragraphs[0]; p.text = "Thank You"
    p.font.size = Pt(50); p.font.bold = True; p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    sub = sl.shapes.add_textbox(Inches(0.8), Inches(4.2), W - Inches(1.6), Inches(0.7))
    q = sub.text_frame.paragraphs[0]; q.text = "Questions?"
    q.font.size = Pt(20); q.font.color.rgb = RGBColor(0x9C, 0xC4, 0xC4); q.alignment = PP_ALIGN.CENTER
    return sl


def build():
    if not SRC.exists():
        print(f"  ! {SRC} not found"); return
    slides = parse(SRC.read_text())
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    for s in slides:
        if s["n"] == 1:
            title_slide(prs, s)
        elif "thank" in s["title"].lower():
            closing_slide(prs, s)
        else:
            content_slide(prs, s, SLIDE_FIGURES.get(s["n"]))
    prs.save(OUT)
    print(f"  deck: {OUT.name}  {len(slides)} slides  ({OUT.stat().st_size//1024} KB)")


if __name__ == "__main__":
    build()
